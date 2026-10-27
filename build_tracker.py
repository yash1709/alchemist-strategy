import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

OUT = sys.argv[1]
N = 1000  # trade rows

wb = Workbook()
dash = wb.active
dash.title = "Dashboard"
st = wb.create_sheet("Settings")
tr = wb.create_sheet("Trades")
gd = wb.create_sheet("Gold")
gu = wb.create_sheet("Guide")

NAVY = "1F3A5F"
HDR = PatternFill("solid", fgColor=NAVY)
INP = PatternFill("solid", fgColor="FFF2CC")
SUB = PatternFill("solid", fgColor="DCE6F1")
GREY = PatternFill("solid", fgColor="F2F2F2")
WHITE_B = Font(bold=True, color="FFFFFF")
B = Font(bold=True)
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def hdr(ws, row, labels, col=1):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=t)
        c.font, c.fill, c.border = WHITE_B, HDR, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=16, color=NAVY)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(italic=True, color="595959")


# ---------------------------------------------------------------- Settings
title(st, "Settings", "Yellow cells are inputs. Everything else is calculated.")
rows = [
    (3, "Total capital (base)", 100000, '#,##0'),
    (4, "Compound profits into capital? (Yes/No)", "No", None),
    (5, "Lump sum, % of capital", 0.06, '0.0%'),
    (6, "SIP, % of capital", 0.016, '0.0%'),
    (7, "Bud buy, % of SIP", 0.5, '0%'),
    (8, "Sell target (gain vs average cost)", 0.05, '0.0%'),
    (9, "Track 2 trigger (fall vs average cost)", -0.05, '0.0%'),
    (10, "Bud mode starts above (% of capital deployed)", 0.8, '0%'),
    (11, "Bud buy trigger (fall vs LAST buy price)", -0.03, '0.0%'),
    (12, "Gold ETF code", "GOLDBEES", None),
]
for r, label, val, fmt in rows:
    st.cell(row=r, column=1, value=label)
    c = st.cell(row=r, column=2, value=val)
    c.fill, c.border = INP, BOX
    if fmt:
        c.number_format = fmt
dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=False)
st.add_data_validation(dv)
dv.add("B4")

calc = [
    (14, "Effective capital", f'=B3+IF(B4="Yes",SUM(Trades!$H$2:$H${N}),0)', '#,##0'),
    (15, "Lump sum amount", "=B14*B5", '#,##0'),
    (16, "SIP amount", "=B14*B6", '#,##0'),
    (17, "Bud buy amount", "=B16*B7", '#,##0'),
    (18, "Capital deployed (open positions)", "=SUM(Dashboard!$G$12:$G$23)", '#,##0'),
    (19, "Deployed % of capital", "=B18/B14", '0.0%'),
    (20, "MODE", '=IF(B19>B10,"BUD","SHIP")', None),
]
for r, label, f, fmt in calc:
    st.cell(row=r, column=1, value=label).font = B
    c = st.cell(row=r, column=2, value=f)
    c.fill, c.border = GREY, BOX
    c.font = B
    if fmt:
        c.number_format = fmt

st["A22"] = "ETF pairs"
st["A22"].font = Font(bold=True, size=12, color=NAVY)
hdr(st, 23, ["Buy day", "Track 1 (Nifty 50)", "Track 2 partner (Bank Nifty)"])
pairs = [(1, "NIFTYBEES", "BANKBEES"), (5, "NIFTYIETF", "BANKIETF"), (10, "SETFNIF50", "SETFNIFBK"),
         (15, "NIFTY1", "BANKNIFTY1"), (20, "HDFCNIFTY", "HDFCNIFBAN"), (25, "NIFTYBETA", "BANKBETA")]
for i, (d, a, b) in enumerate(pairs):
    for j, v in enumerate((d, a, b)):
        c = st.cell(row=24 + i, column=1 + j, value=v)
        c.fill, c.border = INP, BOX
hdr(st, 23, ["All ETF codes (for dropdowns)"], col=5)
for i in range(6):
    st.cell(row=24 + i, column=5, value=f"=B{24+i}")
    st.cell(row=30 + i, column=5, value=f"=C{24+i}")
st.column_dimensions["A"].width = 48
st.column_dimensions["B"].width = 22
st.column_dimensions["C"].width = 28
st.column_dimensions["E"].width = 28

# ---------------------------------------------------------------- Trades
title(tr, "Trade log", "")
tr["A1"] = "Date"
# trades sheet uses row 1 as header (formulas reference rows 2..N)
hdr(tr, 1, ["Date", "ETF code", "Type", "Price", "Qty", "Amount", "Brokerage + charges",
            "Realised profit (SELL rows)", "Notes"])
tr["A1"].font = WHITE_B
tr["A1"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
for r in range(2, N + 1):
    tr[f"F{r}"] = f'=IF(AND(ISNUMBER(D{r}),ISNUMBER(E{r})),D{r}*E{r},"")'
    tr[f"H{r}"] = (
        f'=IF(C{r}="SELL",F{r}-SUMIFS($F$2:$F${N},$B$2:$B${N},B{r},$C$2:$C${N},"<>SELL",'
        f'$A$2:$A${N},">"&_xlfn.MAXIFS($A$2:$A${N},$B$2:$B${N},B{r},$C$2:$C${N},"SELL",$A$2:$A${N},"<"&A{r}),'
        f'$A$2:$A${N},"<="&A{r}),"")'
    )
    tr[f"A{r}"].number_format = "dd-mmm-yyyy"
    for col in "DFGH":
        tr[f"{col}{r}"].number_format = '#,##0.00'
    for col in "ABCDEG":
        tr[f"{col}{r}"].fill = INP
dv1 = DataValidation(type="list", formula1="=Settings!$E$24:$E$35", allow_blank=True)
dv2 = DataValidation(type="list", formula1='"LUMP,SIP,BUD,SELL"', allow_blank=True)
tr.add_data_validation(dv1)
tr.add_data_validation(dv2)
dv1.add(f"B2:B{N}")
dv2.add(f"C2:C{N}")
for col, w in zip("ABCDEFGHI", (14, 16, 10, 12, 10, 14, 14, 18, 36)):
    tr.column_dimensions[col].width = w
tr.freeze_panes = "A2"
tr.row_dimensions[1].height = 32

# ---------------------------------------------------------------- Gold
title(gd, "Gold accumulation", "Log every gold ETF purchase made with booked profits. Gold is never sold.")
hdr(gd, 4, ["Date", "Price", "Qty", "Amount", "Notes"])
for r in range(5, 505):
    gd[f"D{r}"] = f'=IF(AND(ISNUMBER(B{r}),ISNUMBER(C{r})),B{r}*C{r},"")'
    gd[f"A{r}"].number_format = "dd-mmm-yyyy"
    gd[f"B{r}"].number_format = gd[f"D{r}"].number_format = '#,##0.00'
    for col in "ABCE":
        gd[f"{col}{r}"].fill = INP
gsum = [
    (4, "Units held", "=SUM(C5:C504)", '#,##0.00'),
    (5, "Invested", "=SUM(D5:D504)", '#,##0'),
    (6, "Manual price (optional)", None, '#,##0.00'),
    (7, "Price used", '=IF(ISNUMBER(H6),H6,IFERROR(GOOGLEFINANCE("NSE:"&Settings!B12),""))', '#,##0.00'),
    (8, "Current value", '=IF(ISNUMBER(H7),H4*H7,H5)', '#,##0'),
    (9, "Gold P/L", "=H8-H5", '#,##0'),
    (10, "Booked profit not yet put into gold", f"=SUM(Trades!$H$2:$H${N})-H5", '#,##0'),
]
for r, label, f, fmt in gsum:
    gd.cell(row=r, column=7, value=label).font = B
    c = gd.cell(row=r, column=8, value=f)
    c.number_format, c.border = fmt, BOX
    c.fill = INP if f is None else GREY
for col, w in zip("ABCDEFGH", (14, 12, 10, 14, 30, 3, 36, 16)):
    gd.column_dimensions[col].width = w
gd.freeze_panes = "A5"

# ---------------------------------------------------------------- Dashboard
title(dash, "Alchemist 2026 Tracker", "Log trades in the Trades tab; this page tells you what to do today.")
summary_l = [
    (3, "Mode", "=Settings!B20", None),
    (4, "Capital (effective)", "=Settings!B14", '#,##0'),
    (5, "Deployed", "=SUM(G12:G23)", '#,##0'),
    (6, "Deployed %", "=Settings!B19", '0.0%'),
    (7, "Open positions", '=COUNTIF(F12:F23,">0")', '0'),
    (8, "Actions today", '=COUNTIF(R12:R23,"?*")-COUNTIF(R12:R23,"Waits for trigger")', '0'),
]
summary_r = [
    (3, "Realised profit (gross)", f"=SUM(Trades!$H$2:$H${N})", '#,##0'),
    (4, "Unrealised P/L", "=SUM(M12:M23)-SUM(G12:G23)", '#,##0'),
    (5, "Gold value", "=Gold!H8", '#,##0'),
    (6, "Gold P/L", "=Gold!H9", '#,##0'),
    (7, "Total P/L", "=E3+E4+E6", '#,##0'),
    (8, "Total P/L % of base capital", "=E7/Settings!B3", '0.0%'),
    (9, "Profit waiting to go into gold", "=Gold!H10", '#,##0'),
]
for r, label, f, fmt in summary_l:
    dash.cell(row=r, column=1, value=label).font = B
    c = dash.cell(row=r, column=2, value=f)
    c.fill, c.border, c.font = SUB, BOX, B
    if fmt:
        c.number_format = fmt
for r, label, f, fmt in summary_r:
    dash.cell(row=r, column=4, value=label).font = B
    c = dash.cell(row=r, column=5, value=f)
    c.fill, c.border, c.font = SUB, BOX, B
    c.number_format = fmt

H = ["Track", "ETF", "Partner", "Buy day", "Last sell", "Units", "Invested", "Avg cost", "Last buy date",
     "Last buy price", "Cycle start", "Price", "Value", "P/L %", "Sell at", "Fall vs last buy",
     "Next due", "ACTION", "Bud rank", "Manual price", "Last LUMP/SIP", "Bud-eligible fall"]
hdr(dash, 11, H)
dash.row_dimensions[11].height = 32
T, Bc, Cc, Dc, Ac = (f"Trades!$%s$2:$%s${N}" % (x, x) for x in "BCDFA")
Ec = f"Trades!$E$2:$E${N}"


def mif(fn, col, code, extra):
    return f"{fn}(Trades!${col}$2:${col}${N},Trades!$B$2:$B${N},{code}{extra})"


for i in range(12):
    r = 12 + i
    t1 = i < 6
    p = r + 6 if t1 else r - 6
    k = i % 6
    dash[f"A{r}"] = "T1" if t1 else "T2"
    dash[f"B{r}"] = f"=Settings!{'B' if t1 else 'C'}{24+k}"
    dash[f"C{r}"] = f"=B{p}"
    dash[f"D{r}"] = f"=Settings!A{24+k}" if t1 else f'=IF(F{r}>0,DAY(K{r}),"")'
    dash[f"E{r}"] = f'=_xlfn.MAXIFS(Trades!$A$2:$A${N},Trades!$B$2:$B${N},B{r},Trades!$C$2:$C${N},"SELL")'
    dash[f"F{r}"] = (f'=SUMIFS(Trades!$E$2:$E${N},Trades!$B$2:$B${N},B{r},Trades!$C$2:$C${N},"<>SELL")'
                     f'-SUMIFS(Trades!$E$2:$E${N},Trades!$B$2:$B${N},B{r},Trades!$C$2:$C${N},"SELL")')
    dash[f"G{r}"] = (f'=SUMIFS(Trades!$F$2:$F${N},Trades!$B$2:$B${N},B{r},Trades!$C$2:$C${N},"<>SELL",'
                     f'Trades!$A$2:$A${N},">"&E{r})')
    dash[f"H{r}"] = f"=IF(F{r}>0,G{r}/F{r},0)"
    dash[f"I{r}"] = (f'=IF(F{r}>0,_xlfn.MAXIFS(Trades!$A$2:$A${N},Trades!$B$2:$B${N},B{r},Trades!$C$2:$C${N},"<>SELL",'
                     f'Trades!$A$2:$A${N},">"&E{r}),"")')
    dash[f"J{r}"] = (f'=IF(F{r}>0,IFERROR(AVERAGEIFS(Trades!$D$2:$D${N},Trades!$B$2:$B${N},B{r},'
                     f'Trades!$C$2:$C${N},"<>SELL",Trades!$A$2:$A${N},I{r}),0),"")')
    dash[f"K{r}"] = (f'=IF(F{r}>0,_xlfn.MINIFS(Trades!$A$2:$A${N},Trades!$B$2:$B${N},B{r},Trades!$C$2:$C${N},"<>SELL",'
                     f'Trades!$A$2:$A${N},">"&E{r}),"")')
    dash[f"L{r}"] = f'=IF(ISNUMBER(T{r}),T{r},IFERROR(GOOGLEFINANCE("NSE:"&B{r}),""))'
    dash[f"M{r}"] = f'=IF(F{r}>0,IF(ISNUMBER(L{r}),F{r}*L{r},G{r}),0)'
    dash[f"N{r}"] = f'=IF(AND(F{r}>0,ISNUMBER(L{r})),L{r}/H{r}-1,"")'
    dash[f"O{r}"] = f'=IF(F{r}>0,H{r}*(1+Settings!$B$8),"")'
    dash[f"P{r}"] = f'=IF(AND(F{r}>0,ISNUMBER(L{r})),L{r}/J{r}-1,"")'
    dash[f"U{r}"] = (f'=IF(F{r}>0,MAX(_xlfn.MAXIFS(Trades!$A$2:$A${N},Trades!$B$2:$B${N},B{r},Trades!$C$2:$C${N},"LUMP"),'
                     f'_xlfn.MAXIFS(Trades!$A$2:$A${N},Trades!$B$2:$B${N},B{r},Trades!$C$2:$C${N},"SIP")),"")')
    if t1:
        dash[f"Q{r}"] = (
            f'=IF(F{r}>0,DATE(YEAR(U{r}),MONTH(U{r})+1,D{r}),'
            f'IF(E{r}=0,IF(DAY(TODAY())<=D{r},DATE(YEAR(TODAY()),MONTH(TODAY()),D{r}),DATE(YEAR(TODAY()),MONTH(TODAY())+1,D{r})),'
            f'IF(DAY(E{r})<D{r},DATE(YEAR(E{r}),MONTH(E{r}),D{r}),DATE(YEAR(E{r}),MONTH(E{r})+1,D{r}))))'
        )
        dash[f"R{r}"] = (
            f'=IF(F{r}>0,IF(AND(ISNUMBER(L{r}),L{r}>=O{r}),"SELL ALL",'
            f'IF(AND(ISNUMBER(L{r}),L{r}<=H{r}*(1+Settings!$B$9),F{p}=0),"BUY PARTNER - lump "&TEXT(Settings!$B$15,"#,##0"),'
            f'IF(AND(Settings!$B$20="BUD",ISNUMBER(S{r})),"BUD BUY #"&S{r},'
            f'IF(AND(Settings!$B$20="SHIP",TODAY()>=Q{r}),"SIP due - "&TEXT(Settings!$B$16,"#,##0"),"")))),'
            f'IF(AND(Settings!$B$20="SHIP",TODAY()>=Q{r}),"LUMP due - "&TEXT(Settings!$B$15,"#,##0"),""))'
        )
    else:
        dash[f"Q{r}"] = f'=IF(F{r}>0,DATE(YEAR(U{r}),MONTH(U{r})+1,DAY(K{r})),"")'
        dash[f"R{r}"] = (
            f'=IF(F{r}>0,IF(AND(ISNUMBER(L{r}),L{r}>=O{r}),"SELL ALL",'
            f'IF(AND(Settings!$B$20="BUD",ISNUMBER(S{r})),"BUD BUY #"&S{r},'
            f'IF(AND(Settings!$B$20="SHIP",TODAY()>=Q{r}),"SIP due - "&TEXT(Settings!$B$16,"#,##0"),""))),'
            f'"Waits for trigger")'
        )
    dash[f"V{r}"] = (f'=IF(AND(Settings!$B$20="BUD",F{r}>0,ISNUMBER(P{r})),IF(P{r}<=Settings!$B$11,P{r},""),"")')
    dash[f"S{r}"] = f'=IF(ISNUMBER(V{r}),RANK(V{r},$V$12:$V$23,1),"")'
    for col in "EIKQ":
        dash[f"{col}{r}"].number_format = "dd-mmm-yy;;"
    for col in "GM":
        dash[f"{col}{r}"].number_format = '#,##0'
    for col in "HJLOT":
        dash[f"{col}{r}"].number_format = '#,##0.00'
    for col in "NPV":
        dash[f"{col}{r}"].number_format = '0.0%'
    dash[f"T{r}"].fill = INP
    dash[f"R{r}"].font = B
    for col in range(1, 23):
        dash.cell(row=r, column=col).border = BOX
    if not t1:
        for col in range(1, 23):
            if col != 20:
                dash.cell(row=r, column=col).fill = GREY

rng = "R12:R23"
dash.conditional_formatting.add(rng, FormulaRule(formula=['LEFT($R12,4)="SELL"'], fill=PatternFill("solid", bgColor="C6EFCE"), font=Font(bold=True, color="006100")))
dash.conditional_formatting.add(rng, FormulaRule(formula=['OR(LEFT($R12,3)="BUY",LEFT($R12,3)="BUD")'], fill=PatternFill("solid", bgColor="FFC7CE"), font=Font(bold=True, color="9C0006")))
dash.conditional_formatting.add(rng, FormulaRule(formula=['OR(LEFT($R12,3)="SIP",LEFT($R12,4)="LUMP")'], fill=PatternFill("solid", bgColor="FFEB9C"), font=Font(bold=True, color="9C5700")))
dash.conditional_formatting.add("B3", FormulaRule(formula=['$B$3="BUD"'], fill=PatternFill("solid", bgColor="FFC7CE")))
dash.conditional_formatting.add("N12:N23", FormulaRule(formula=['AND(ISNUMBER($N12),$N12<0)'], font=Font(color="C00000")))

for col, w in zip("ABCDEFGHIJKLMNOPQRSTUV", (22, 14, 14, 9, 11, 9, 11, 10, 12, 11, 12, 10, 11, 8, 10, 11, 12, 30, 9, 12, 12, 12)):
    dash.column_dimensions[col].width = w
dash.freeze_panes = "C12"
dash["A10"] = "Positions (T1 = Nifty 50 ETF, T2 = Bank Nifty partner)"
dash["A10"].font = Font(bold=True, size=12, color=NAVY)

# ---------------------------------------------------------------- Guide
title(gu, "How to use this tracker")
lines = [
    "SETUP",
    "1. Settings tab: enter your total capital (yellow). Lump sum, SIP and Bud amounts calculate themselves.",
    "2. Confirm the six ETF pairs and buy days (yellow). Change the Gold ETF code if you use another.",
    "",
    "LOGGING (Trades tab - yellow cells only; add rows in date order)",
    "- LUMP: first buy of an ETF cycle.   SIP: monthly add.   BUD: a Bud-mode dip buy.",
    "- SELL: always enter the FULL quantity of that ETF. The sheet computes profit for the cycle.",
    "- Record the real broker price and quantity. Brokerage is optional.",
    "- After a SELL, buy gold with the profit and log it in the Gold tab (Dashboard shows the amount waiting).",
    "",
    "DAILY (after market, 2 minutes)",
    "- Open Dashboard. Anything in the ACTION column is something to do tomorrow/now:",
    "   SELL ALL (green): price reached average cost +5%.",
    "   BUY PARTNER (red): the Nifty ETF is 5% under average cost; lump-sum its Bank Nifty partner.",
    "   BUD BUY #n (red): Bud mode is on; buy ONLY the #1 ETF today (one Bud buy per day).",
    "   SIP due / LUMP due (amber): scheduled monthly buy. If it falls on a holiday, buy next trading day.",
    "",
    "RULES BUILT IN",
    "- Mode flips to BUD when deployed capital is above 80% of capital; SIPs and lump-sum restarts pause in BUD mode.",
    "- Next due date for a restarted ETF: this month if sold before its buy day, otherwise next month.",
    "- Track 2 partners wait for a trigger; their SIP day is the day of their first lump sum.",
    "",
    "PRICES",
    "- Live prices use GOOGLEFINANCE (Google Sheets only). Type a price in 'Manual price' to override or if a quote is missing.",
    "- Opening this file in Excel works too, but live prices need manual entry.",
    "",
    "If capital reaches 100%, lower the sell target in Settings (5% -> 4% -> 3%) to free capital, then restore it.",
    "Not investment advice. Past backtests are not a promise.",
]
for i, t in enumerate(lines):
    c = gu.cell(row=3 + i, column=1, value=t)
    if t.isupper() or t in ("SETUP",):
        c.font = Font(bold=True, color=NAVY)
gu.column_dimensions["A"].width = 120

wb.save(OUT)
print("saved", OUT)
